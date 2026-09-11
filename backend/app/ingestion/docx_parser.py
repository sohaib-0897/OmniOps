import docx
from typing import List, Dict, Any
from app.core.config import settings

def parse_docx_document(file_path: str) -> List[Dict[str, Any]]:
    """Extract paragraphs, headings, and tables from DOCX documents."""
    doc = docx.Document(file_path)
    if len(doc.paragraphs) > settings.MAX_DOCX_PARAGRAPHS:
        raise ValueError("DOCX paragraph limit exceeded.")
    chunks = []
    
    current_heading = "Document Root"
    current_paragraphs = []
    
    extracted_chars = 0
    for para in doc.paragraphs:
        text = para.text.strip()
        extracted_chars += len(text)
        if extracted_chars > settings.MAX_EXTRACTED_TEXT_CHARS:
            raise ValueError("DOCX extracted text limit exceeded.")
        if not text:
            continue
            
        if para.style.name.startswith("Heading"):
            if current_paragraphs:
                chunks.append({
                    "content": f"## {current_heading}\n" + "\n\n".join(current_paragraphs),
                    "modality": "docx",
                    "page_number": None,
                    "extraction_method": "DOCX_TEXT",
                    "metadata": {"heading": current_heading, "extraction_method": "DOCX_TEXT"}
                })
                current_paragraphs = []
            current_heading = text
        else:
            current_paragraphs.append(text)
            
    if current_paragraphs:
        chunks.append({
            "content": f"## {current_heading}\n" + "\n\n".join(current_paragraphs),
            "modality": "docx",
            "page_number": None,
            "extraction_method": "DOCX_TEXT",
            "metadata": {"heading": current_heading, "extraction_method": "DOCX_TEXT"}
        })
        
    # Extract tables
    for t_idx, table in enumerate(doc.tables):
        rows_data = []
        for row in table.rows:
            row_vals = [cell.text.strip() for cell in row.cells]
            extracted_chars += sum(len(value) for value in row_vals)
            if extracted_chars > settings.MAX_EXTRACTED_TEXT_CHARS:
                raise ValueError("DOCX extracted text limit exceeded.")
            rows_data.append(" | ".join(row_vals))
        if rows_data:
            table_md = "\n".join(rows_data)
            chunks.append({
                "content": f"### Table {t_idx + 1}\n" + table_md,
                "modality": "docx",
                "page_number": None,
                "extraction_method": "TABLE_EXTRACTION",
                "metadata": {"table_index": t_idx + 1, "extraction_method": "TABLE_EXTRACTION"}
            })
            
    return chunks
